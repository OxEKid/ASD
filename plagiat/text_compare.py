"""Модуль 4. Сравнение текстовых отчётов.
Ответственный: Каныгин К. К.

Собственный конвейер обработки отчётов:
  1. извлечение текста из DOCX (python-docx), PDF (PyMuPDF) и TXT;
  2. нормализация: один регистр, без пунктуации и лишних пробелов;
  3. исключение шаблонных частей: абзацы из файлов шаблонов (титульный лист,
     текст задания) и абзацы, которые встречаются у большинства работ;
  4. разбиение на шинглы (перекрывающиеся фрагменты из n слов);
  5. отпечатки шинглов через hashlib и сравнение наборов отпечатков.
"""
from __future__ import annotations
import hashlib, io, re, zipfile
from collections import Counter
from xml.etree import ElementTree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def extract_text(name: str | None, data: bytes | None) -> str:
    if not name or data is None:
        return ""
    low = name.lower()
    try:
        if low.endswith(".docx"):
            try:
                import docx  # python-docx
                d = docx.Document(io.BytesIO(data))
                parts = [p.text for p in d.paragraphs]
                for t in d.tables:
                    for row in t.rows:
                        parts.append(" ".join(c.text for c in row.cells))
                return "\n".join(p for p in parts if p.strip())
            except ImportError:
                root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
                paras = ["".join(t.text or "" for t in p.iter(W + "t")) for p in root.iter(W + "p")]
                return "\n".join(p for p in paras if p.strip())
        if low.endswith(".pdf"):
            try:
                import pymupdf
            except ImportError:
                try:
                    import fitz as pymupdf
                except ImportError:
                    from pypdf import PdfReader
                    return "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(data)).pages)
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc)
        if low.endswith(".txt"):
            for enc in ("utf-8-sig", "cp1251"):
                try:
                    return data.decode(enc)
                except UnicodeDecodeError:
                    pass
    except Exception:
        return ""
    return ""

def _norm(s: str) -> str:
    return " ".join(re.findall(r"[a-zа-яё0-9]+", s.lower().replace("ё", "е")))

def paragraphs(text: str) -> list[str]:
    return [p for p in (_norm(x) for x in re.split(r"\n+", text)) if p]

def find_templates(all_texts: list[str], share: float = 0.5, template_texts: list[str] = ()) -> set[str]:
    """Шаблонные абзацы: из файлов шаблонов и те, что встречаются у доли работ >= share."""
    res = {p for t in template_texts for p in paragraphs(t)}
    if len(all_texts) >= 3:
        cnt = Counter(p for t in all_texts for p in set(paragraphs(t)))
        res |= {p for p, c in cnt.items() if c / len(all_texts) >= share}
    return res

def _fp(sh: tuple) -> int:
    return int.from_bytes(hashlib.md5(" ".join(sh).encode()).digest()[:8], "big")

def words(text: str, templates: set[str]) -> list[str]:
    return [w for p in paragraphs(text) if p not in templates for w in p.split()]

def compare(wa: list[str], wb: list[str], n: int = 4) -> dict:
    if len(wa) < n or len(wb) < n:
        return {"similarity": None, "fragments": []}
    ha = [_fp(tuple(wa[i:i+n])) for i in range(len(wa) - n + 1)]   # отпечатки шинглов
    hb = [_fp(tuple(wb[i:i+n])) for i in range(len(wb) - n + 1)]
    sa, sb = set(ha), set(hb)
    common = sa & sb
    sim = len(common) / min(len(sa), len(sb))
    # склеиваем подряд идущие общие шинглы в фрагменты (по тексту A)
    frags, cur = [], []
    for i in range(len(wa) - n + 1):
        if ha[i] in common:
            cur = cur + [wa[i+n-1]] if cur else list(wa[i:i+n])
        elif cur:
            frags.append(" ".join(cur)); cur = []
    if cur:
        frags.append(" ".join(cur))
    frags = sorted((f for f in frags if len(f.split()) >= 6), key=len, reverse=True)[:8]
    return {"similarity": round(min(sim, 1.0), 4), "fragments": frags}
